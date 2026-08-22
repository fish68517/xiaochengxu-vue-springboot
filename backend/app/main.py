import csv
import io
from datetime import date, datetime
from uuid import uuid4

from fastapi import Depends, FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session, selectinload

from .adapters import notification_service, storage_service
from .config import settings
from .database import get_db
from .models import (
    Bill, BillItem, Building, FeeItem, House, ImportBatch, MaintenanceRecord,
    Notice, NotificationOutbox, Payment, Receipt, RepairOrder, Renovation, User,
)
from .security import create_access_token, get_current_user, hash_password, require_roles, verify_password


app = FastAPI(title="加装电梯后期管理平台 API", version="0.2.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=settings.cors_list,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)
app.mount("/uploads", StaticFiles(directory=storage_service.root), name="uploads")


def response(data=None, message="操作成功"):
    return {"code": 0, "message": message, "data": data}


def money(fen: int) -> str:
    return f"{fen / 100:.2f}"


def house_label(house: House) -> str:
    return f"{house.building.community_name}{house.building.name} {house.room_no}"


def session_data(db: Session, user: User) -> dict:
    house = None
    if user.house_id:
        house = db.scalar(select(House).options(selectinload(House.building)).where(House.id == user.house_id))
    return {
        "id": user.id,
        "username": user.username,
        "displayName": user.display_name,
        "role": user.role,
        "phone": user.phone,
        "enabled": user.enabled,
        "house": {"id": house.id, "displayName": house_label(house), "buildingId": house.building_id} if house else None,
    }


def bill_data(bill: Bill) -> dict:
    effective_status = "OVERDUE" if bill.status == "PENDING" and bill.due_date < date.today() else bill.status
    return {
        "id": bill.id,
        "billNo": bill.bill_no,
        "houseId": bill.house_id,
        "houseDisplayName": house_label(bill.house),
        "title": bill.title,
        "billingPeriod": bill.billing_period,
        "totalAmountFen": bill.total_amount_fen,
        "totalAmount": money(bill.total_amount_fen),
        "paidAmountFen": bill.paid_amount_fen,
        "dueDate": bill.due_date.isoformat(),
        "status": effective_status,
        "generatedBy": bill.generated_by,
        "createdAt": bill.created_at.isoformat(),
        "paidAt": bill.paid_at.isoformat() if bill.paid_at else None,
        "items": [{
            "id": item.id,
            "feeItemCode": item.fee_item_code,
            "feeItemName": item.fee_item_name,
            "amountFen": item.amount_fen,
            "amount": money(item.amount_fen),
            "remark": item.remark,
        } for item in bill.items],
    }


def repair_data(item: RepairOrder) -> dict:
    return {
        "id": item.id,
        "orderNo": item.order_no,
        "houseId": item.house_id,
        "house": house_label(item.house),
        "description": item.description,
        "imageUrls": item.image_urls,
        "status": item.status,
        "contactName": item.contact_name,
        "contactPhone": item.contact_phone,
        "progressText": item.progress_text,
        "resultText": item.result_text,
        "technicianId": item.technician_id,
        "technicianName": item.technician.display_name if item.technician else None,
        "createdAt": item.created_at.isoformat(),
        "updatedAt": item.updated_at.isoformat(),
    }


def maintenance_data(item: MaintenanceRecord) -> dict:
    return {
        "id": item.id, "title": item.title, "buildingId": item.building_id,
        "elevatorName": item.elevator_name, "maintenanceDate": item.maintenance_date.isoformat(),
        "pdfUrl": item.pdf_url, "fileSize": item.file_size, "uploaderName": item.uploader_name,
        "createdAt": item.created_at.isoformat(),
    }


def notice_data(item: Notice) -> dict:
    return {
        "id": item.id, "title": item.title, "category": item.category,
        "content": item.content, "pinned": item.pinned, "publishedAt": item.published_at.isoformat(),
    }


def renovation_data(item: Renovation) -> dict:
    return {
        "id": item.id, "houseId": item.house_id, "applicantName": item.applicant_name,
        "phone": item.phone, "plan": item.plan, "status": item.status,
        "createdAt": item.created_at.isoformat(),
    }


def payment_data(payment: Payment) -> dict:
    return {
        "id": payment.id, "paymentNo": payment.payment_no, "billId": payment.bill_id,
        "billNo": payment.bill.bill_no if payment.bill else None, "amountFen": payment.amount_fen,
        "amount": money(payment.amount_fen), "status": payment.status, "provider": payment.provider,
        "providerTransactionId": payment.provider_transaction_id, "createdAt": payment.created_at.isoformat(),
        "paidAt": payment.paid_at.isoformat() if payment.paid_at else None,
    }


def receipt_data(item: Receipt) -> dict:
    return {
        "id": item.id, "applicationNo": item.application_no, "billId": item.bill_id,
        "paymentId": item.payment_id, "applicantName": item.applicant_name, "phone": item.phone,
        "title": item.title, "remark": item.remark, "status": item.status,
        "createdAt": item.created_at.isoformat(),
    }


class LoginRequest(BaseModel):
    username: str = Field(min_length=2, max_length=64)
    password: str = Field(min_length=8, max_length=128)


@app.get("/api/health")
def health(db: Session = Depends(get_db)):
    db.execute(select(1))
    return response({
        "status": "ok", "appEnv": settings.app_env, "database": "mysql",
        "authMode": settings.auth_mode, "paymentMode": settings.payment_mode,
        "storageBackend": settings.storage_backend, "notificationMode": settings.notification_mode,
    })


@app.post("/api/auth/login")
def login(payload: LoginRequest, db: Session = Depends(get_db)):
    user = db.scalar(select(User).where(User.username == payload.username))
    if not user or not user.enabled or not verify_password(payload.password, user.password_hash):
        raise HTTPException(401, "用户名或密码错误")
    return response({"token": create_access_token(user), "user": session_data(db, user)}, "登录成功")


@app.get("/api/auth/me")
def auth_me(user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    return response(session_data(db, user))


@app.post("/api/auth/logout")
def logout(user: User = Depends(get_current_user)):
    return response({"username": user.username}, "已退出登录")


@app.get("/api/bills/my")
def my_bills(status: str | None = None, user: User = Depends(require_roles("OWNER")), db: Session = Depends(get_db)):
    if not user.house_id:
        raise HTTPException(409, "当前账号未绑定房屋")
    query = select(Bill).options(selectinload(Bill.house).selectinload(House.building), selectinload(Bill.items)).where(Bill.house_id == user.house_id).order_by(Bill.created_at.desc())
    items = [bill_data(item) for item in db.scalars(query).all()]
    if status:
        items = [item for item in items if item["status"] == status]
    return response(items)


@app.get("/api/bills/{bill_id}")
def bill_detail(bill_id: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    bill = db.scalar(select(Bill).options(selectinload(Bill.house).selectinload(House.building), selectinload(Bill.items)).where(Bill.id == bill_id))
    if not bill:
        raise HTTPException(404, "账单不存在")
    if user.role == "OWNER" and bill.house_id != user.house_id:
        raise HTTPException(403, "无权查看其他房屋账单")
    if user.role not in ("OWNER", "ADMIN"):
        raise HTTPException(403, "当前账号无权查看账单")
    return response(bill_data(bill))


class PaymentRequest(BaseModel):
    billId: int


@app.post("/api/payments/manual-request")
def request_manual_payment(payload: PaymentRequest, user: User = Depends(require_roles("OWNER")), db: Session = Depends(get_db)):
    bill = db.scalar(select(Bill).where(Bill.id == payload.billId).with_for_update())
    if not bill:
        raise HTTPException(404, "账单不存在")
    if bill.house_id != user.house_id:
        raise HTTPException(403, "无权操作此账单")
    if bill.status == "PAID":
        existing = db.scalar(select(Payment).where(Payment.bill_id == bill.id, Payment.status == "SUCCESS"))
        return response({"status": "SUCCESS", "paymentId": existing.id if existing else None}, "账单已支付")
    existing = db.scalar(select(Payment).where(Payment.bill_id == bill.id, Payment.provider == "MANUAL", Payment.status == "PENDING"))
    if existing:
        return response({"status": existing.status, "paymentId": existing.id, "idempotent": True}, "已提交确认申请")
    payment = Payment(payment_no=f"MAN{datetime.now():%Y%m%d%H%M%S}{uuid4().hex[:6]}", bill_id=bill.id, amount_fen=bill.total_amount_fen, provider="MANUAL", status="PENDING")
    db.add(payment)
    db.flush()
    notification_service.send(db, "MANUAL_PAYMENT_REQUESTED", "role:ADMIN", {"paymentId": payment.id})
    db.commit()
    return response({"status": payment.status, "paymentId": payment.id}, "已提交线下付款确认申请")


@app.get("/api/payments/my")
def my_payments(user: User = Depends(require_roles("OWNER")), db: Session = Depends(get_db)):
    items = db.scalars(select(Payment).options(selectinload(Payment.bill)).join(Bill).where(Bill.house_id == user.house_id).order_by(Payment.created_at.desc())).all()
    return response([payment_data(item) for item in items])


class ReceiptRequest(BaseModel):
    billId: int
    applicantName: str = Field(min_length=1, max_length=64)
    phone: str = Field(min_length=6, max_length=24)
    title: str = Field(default="", max_length=100)
    remark: str = Field(default="", max_length=255)


@app.post("/api/receipts")
def create_receipt(payload: ReceiptRequest, user: User = Depends(require_roles("OWNER")), db: Session = Depends(get_db)):
    bill = db.scalar(select(Bill).where(Bill.id == payload.billId, Bill.house_id == user.house_id))
    if not bill or bill.status != "PAID":
        raise HTTPException(409, "仅已缴账单可以申请收据")
    payment = db.scalar(select(Payment).where(Payment.bill_id == bill.id, Payment.status == "SUCCESS"))
    if not payment:
        raise HTTPException(409, "未找到成功缴费记录")
    existing = db.scalar(select(Receipt).where(Receipt.bill_id == bill.id))
    if existing:
        return response({"id": existing.id, "status": existing.status}, "已提交过收据申请")
    item = Receipt(application_no=f"RC{datetime.now():%Y%m%d%H%M%S}{uuid4().hex[:4]}", bill_id=bill.id, payment_id=payment.id, applicant_name=payload.applicantName, phone=payload.phone, title=payload.title, remark=payload.remark)
    db.add(item)
    notification_service.send(db, "RECEIPT_APPLIED", "role:ADMIN", {"billId": bill.id})
    db.commit()
    return response({"id": item.id, "status": item.status})


@app.get("/api/receipts/my")
def my_receipts(user: User = Depends(require_roles("OWNER")), db: Session = Depends(get_db)):
    items = db.scalars(select(Receipt).join(Bill).where(Bill.house_id == user.house_id).order_by(Receipt.created_at.desc())).all()
    return response([receipt_data(item) for item in items])


class RepairRequest(BaseModel):
    description: str = Field(min_length=5)
    contactName: str = Field(min_length=1, max_length=64)
    contactPhone: str = Field(min_length=6, max_length=24)
    imageUrls: list[str] = []


@app.get("/api/repairs/my")
def my_repairs(user: User = Depends(require_roles("OWNER")), db: Session = Depends(get_db)):
    items = db.scalars(select(RepairOrder).options(selectinload(RepairOrder.house).selectinload(House.building), selectinload(RepairOrder.technician)).where(RepairOrder.house_id == user.house_id).order_by(RepairOrder.created_at.desc())).all()
    return response([repair_data(item) for item in items])


@app.post("/api/repairs")
def create_repair(payload: RepairRequest, user: User = Depends(require_roles("OWNER")), db: Session = Depends(get_db)):
    if not user.house_id:
        raise HTTPException(409, "请先绑定房屋")
    item = RepairOrder(order_no=f"BX{datetime.now():%Y%m%d%H%M%S}{uuid4().hex[:4]}", house_id=user.house_id, description=payload.description, image_urls=payload.imageUrls, contact_name=payload.contactName, contact_phone=payload.contactPhone)
    db.add(item)
    db.flush()
    notification_service.send(db, "REPAIR_CREATED", "role:ADMIN", {"orderNo": item.order_no})
    db.commit()
    return response({"id": item.id, "orderNo": item.order_no})


@app.get("/api/technician/repairs")
def technician_repairs(user: User = Depends(require_roles("TECHNICIAN")), db: Session = Depends(get_db)):
    items = db.scalars(select(RepairOrder).options(selectinload(RepairOrder.house).selectinload(House.building), selectinload(RepairOrder.technician)).where(or_(RepairOrder.technician_id == user.id, RepairOrder.technician_id.is_(None))).order_by(RepairOrder.created_at.desc())).all()
    return response([repair_data(item) for item in items])


class RepairAction(BaseModel):
    action: str
    text: str = Field(default="", max_length=255)


@app.post("/api/technician/repairs/{repair_id}/action")
def technician_action(repair_id: int, payload: RepairAction, user: User = Depends(require_roles("TECHNICIAN")), db: Session = Depends(get_db)):
    item = db.scalar(select(RepairOrder).where(RepairOrder.id == repair_id).with_for_update())
    if not item:
        raise HTTPException(404, "工单不存在")
    if item.technician_id not in (None, user.id):
        raise HTTPException(403, "该工单未派给当前师傅")
    item.technician_id = user.id
    if payload.action in ("accept", "progress"):
        item.status = "REPAIRING"
        item.progress_text = payload.text or "维修师傅已接单，正在处理中"
    elif payload.action == "complete":
        item.status = "COMPLETED"
        item.result_text = payload.text or "故障已处理，电梯恢复正常运行"
    else:
        raise HTTPException(400, "不支持的工单动作")
    notification_service.send(db, "REPAIR_STATUS_CHANGED", f"house:{item.house_id}", {"orderNo": item.order_no, "status": item.status})
    db.commit()
    return response({"id": item.id, "status": item.status})


@app.get("/api/maintenance-records")
def maintenance_records(user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    query = select(MaintenanceRecord).order_by(MaintenanceRecord.maintenance_date.desc())
    if user.role == "OWNER":
        if not user.building_id:
            raise HTTPException(403, "未绑定楼栋")
        query = query.where(MaintenanceRecord.building_id == user.building_id)
    items = db.scalars(query).all()
    return response([maintenance_data(item) for item in items])


@app.get("/api/notices")
def notices(user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    items = db.scalars(select(Notice).order_by(Notice.pinned.desc(), Notice.published_at.desc())).all()
    return response([notice_data(item) for item in items])


@app.get("/api/renovations/my")
def my_renovations(user: User = Depends(require_roles("OWNER")), db: Session = Depends(get_db)):
    items = db.scalars(select(Renovation).where(Renovation.house_id == user.house_id).order_by(Renovation.created_at.desc())).all()
    return response([renovation_data(item) for item in items])


class RenovationRequest(BaseModel):
    applicantName: str = Field(min_length=1, max_length=64)
    phone: str = Field(min_length=6, max_length=24)
    plan: str = Field(min_length=2, max_length=255)


@app.post("/api/renovations")
def create_renovation(payload: RenovationRequest, user: User = Depends(require_roles("OWNER")), db: Session = Depends(get_db)):
    if not user.house_id:
        raise HTTPException(409, "请先绑定房屋")
    item = Renovation(house_id=user.house_id, applicant_name=payload.applicantName, phone=payload.phone, plan=payload.plan, status="DEPOSIT_UNPAID")
    db.add(item)
    db.flush()
    notification_service.send(db, "RENOVATION_CREATED", "role:ADMIN", {"renovationId": item.id})
    db.commit()
    return response({"id": item.id, "status": item.status})


@app.post("/api/uploads")
async def upload(file: UploadFile = File(...), folder: str = "common", user: User = Depends(get_current_user)):
    result = await storage_service.save(file, folder)
    return response(result)


@app.get("/api/admin/dashboard")
def admin_dashboard(admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    pending = db.scalar(select(func.count()).select_from(RepairOrder).where(RepairOrder.status == "PENDING")) or 0
    overdue_condition = or_(Bill.status == "OVERDUE", (Bill.status == "PENDING") & (Bill.due_date < date.today()))
    arrears = db.scalar(select(func.count()).select_from(Bill).where(overdue_condition)) or 0
    receivable = db.scalar(select(func.sum(Bill.total_amount_fen - Bill.paid_amount_fen)).where(Bill.status.in_(["PENDING", "OVERDUE"]))) or 0
    paid = db.scalar(select(func.sum(Bill.paid_amount_fen))) or 0
    return response({
        "pendingRepairs": pending, "arrearsHouseholds": arrears,
        "receivable": money(receivable), "paid": money(paid),
        "maintenanceCount": db.scalar(select(func.count()).select_from(MaintenanceRecord)) or 0,
        "noticeCount": db.scalar(select(func.count()).select_from(Notice)) or 0,
    })


@app.get("/api/admin/fee-items")
def admin_fee_items(admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    items = db.scalars(select(FeeItem).order_by(FeeItem.sort_order)).all()
    return response([{
        "id": item.id, "code": item.code, "name": item.name, "description": item.description,
        "defaultAmountFen": item.default_amount_fen, "defaultAmount": money(item.default_amount_fen or 0),
        "enabled": item.enabled, "sortOrder": item.sort_order,
    } for item in items])


class FeeItemPayload(BaseModel):
    code: str = Field(min_length=2, max_length=32)
    name: str = Field(min_length=1, max_length=100)
    description: str = Field(default="", max_length=255)
    defaultAmountFen: int | None = Field(default=None, ge=0)
    enabled: bool = True
    sortOrder: int = 0


@app.post("/api/admin/fee-items")
def create_fee_item(payload: FeeItemPayload, admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    if db.scalar(select(FeeItem.id).where(FeeItem.code == payload.code)):
        raise HTTPException(409, "收费项目编码已存在")
    item = FeeItem(code=payload.code, name=payload.name, description=payload.description, default_amount_fen=payload.defaultAmountFen, enabled=payload.enabled, sort_order=payload.sortOrder)
    db.add(item)
    db.commit()
    return response({"id": item.id})


@app.put("/api/admin/fee-items/{item_id}")
def update_fee_item(item_id: int, payload: FeeItemPayload, admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    item = db.get(FeeItem, item_id)
    if not item:
        raise HTTPException(404, "收费项目不存在")
    if db.scalar(select(FeeItem.id).where(FeeItem.code == payload.code, FeeItem.id != item_id)):
        raise HTTPException(409, "收费项目编码已存在")
    item.code, item.name, item.description = payload.code, payload.name, payload.description
    item.default_amount_fen, item.enabled, item.sort_order = payload.defaultAmountFen, payload.enabled, payload.sortOrder
    db.commit()
    return response({"id": item.id})


@app.get("/api/admin/bills")
def admin_bills(admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    items = db.scalars(select(Bill).options(selectinload(Bill.house).selectinload(House.building), selectinload(Bill.items)).order_by(Bill.created_at.desc())).all()
    return response([bill_data(item) for item in items])


class BillItemPayload(BaseModel):
    feeItemId: int
    amountFen: int = Field(ge=0)
    remark: str = Field(default="", max_length=255)


class BillPayload(BaseModel):
    houseId: int
    title: str = Field(min_length=2, max_length=180)
    billingPeriod: str = Field(min_length=4, max_length=32)
    dueDate: date
    items: list[BillItemPayload] = Field(min_length=1)


@app.post("/api/admin/bills")
def create_bill(payload: BillPayload, admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    house = db.get(House, payload.houseId)
    if not house:
        raise HTTPException(404, "房屋不存在")
    if db.scalar(select(Bill.id).where(Bill.house_id == payload.houseId, Bill.billing_period == payload.billingPeriod, Bill.title == payload.title)):
        raise HTTPException(409, "该房屋、周期和标题的账单已存在")
    fee_ids = {item.feeItemId for item in payload.items}
    fees = {item.id: item for item in db.scalars(select(FeeItem).where(FeeItem.id.in_(fee_ids), FeeItem.enabled.is_(True))).all()}
    if fee_ids != set(fees):
        raise HTTPException(400, "包含不存在或已停用的收费项目")
    bill = Bill(bill_no=f"BILL-{payload.houseId}-{payload.billingPeriod}-{uuid4().hex[:6]}", house_id=payload.houseId, title=payload.title, billing_period=payload.billingPeriod, due_date=payload.dueDate, status="PENDING", generated_by="MANUAL", total_amount_fen=sum(item.amountFen for item in payload.items))
    db.add(bill)
    db.flush()
    for row in payload.items:
        fee = fees[row.feeItemId]
        bill.items.append(BillItem(bill_id=bill.id, fee_item_id=fee.id, fee_item_code=fee.code, fee_item_name=fee.name, amount_fen=row.amountFen, remark=row.remark))
    notification_service.send(db, "BILL_CREATED", f"house:{house.id}", {"billId": bill.id})
    db.commit()
    return response({"id": bill.id, "billNo": bill.bill_no})


def parse_upload(filename: str, content: bytes) -> list[dict]:
    required = ["communityName", "buildingName", "roomNo", "residentCode", "billingPeriod", "billTitle", "feeItemCode", "feeItemName", "amountYuan", "dueDate"]
    if filename.lower().endswith(".csv"):
        rows = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
    elif filename.lower().endswith((".xlsx", ".xlsm")):
        values = list(load_workbook(io.BytesIO(content), read_only=True, data_only=True).active.iter_rows(values_only=True))
        if not values:
            raise HTTPException(400, "导入文件没有数据")
        headers = [str(value or "") for value in values[0]]
        rows = [dict(zip(headers, row)) for row in values[1:]]
    else:
        raise HTTPException(400, "仅支持 CSV 或 XLSX")
    if not rows:
        raise HTTPException(400, "导入文件没有数据")
    missing = [field for field in required if field not in rows[0]]
    if missing:
        raise HTTPException(400, f"缺少字段：{', '.join(missing)}")
    return rows


@app.post("/api/admin/billing-imports/parse")
async def parse_billing_import(file: UploadFile = File(...), admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    raw_rows = parse_upload(file.filename or "", await file.read())
    preview, seen = [], set()
    for index, row in enumerate(raw_rows, 2):
        errors = []
        house = db.scalar(select(House).where(House.room_no == str(row.get("roomNo", "")).strip(), House.resident_code == str(row.get("residentCode", "")).strip()))
        fee = db.scalar(select(FeeItem).where(FeeItem.code == str(row.get("feeItemCode", "")).strip(), FeeItem.enabled.is_(True)))
        try:
            amount_fen = int(round(float(row.get("amountYuan")) * 100))
            if amount_fen < 0:
                errors.append("金额不能小于0")
        except (TypeError, ValueError):
            amount_fen = 0
            errors.append("金额格式错误")
        try:
            due = date.fromisoformat(str(row.get("dueDate"))[:10])
        except ValueError:
            due = None
            errors.append("截止日期格式错误")
        if not house:
            errors.append("房号不存在或住户编码不匹配")
        if not fee:
            errors.append("费用项目不存在或未启用")
        key = (row.get("buildingName"), row.get("roomNo"), row.get("billingPeriod"), row.get("billTitle"), row.get("feeItemCode"))
        if key in seen:
            errors.append("文件内重复费用项")
        seen.add(key)
        preview.append({"rowNo": index, **{field: str(row.get(field, "") or "") for field in row}, "amountFen": amount_fen, "dueDate": due.isoformat() if due else "", "houseId": house.id if house else None, "feeItemId": fee.id if fee else None, "valid": not errors, "errors": errors})
    batch = ImportBatch(batch_no=f"IMP{datetime.now():%Y%m%d%H%M%S}{uuid4().hex[:4]}", filename=file.filename or "", rows_json=preview)
    db.add(batch)
    db.commit()
    return response({"batchId": batch.id, "batchNo": batch.batch_no, "validCount": sum(1 for row in preview if row["valid"]), "invalidCount": sum(1 for row in preview if not row["valid"]), "rows": preview})


@app.post("/api/admin/billing-imports/{batch_id}/commit")
def commit_billing_import(batch_id: int, admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    batch = db.scalar(select(ImportBatch).where(ImportBatch.id == batch_id).with_for_update())
    if not batch:
        raise HTTPException(404, "导入批次不存在")
    if batch.status == "COMMITTED":
        return response({"batchId": batch.id, "idempotent": True}, "该批次已经提交")
    groups: dict[tuple, list[dict]] = {}
    for row in batch.rows_json:
        if row["valid"]:
            groups.setdefault((row["houseId"], row["billingPeriod"], row["billTitle"], row["dueDate"]), []).append(row)
    created = skipped = 0
    for (house_id, period, title, due), rows in groups.items():
        if db.scalar(select(Bill).where(Bill.house_id == house_id, Bill.billing_period == period, Bill.title == title)):
            skipped += 1
            continue
        bill = Bill(bill_no=f"BILL-{house_id}-{period}-{uuid4().hex[:5]}", house_id=house_id, title=title, billing_period=period, due_date=date.fromisoformat(due), status="PENDING", generated_by="IMPORT", total_amount_fen=sum(row["amountFen"] for row in rows))
        db.add(bill)
        db.flush()
        for row in rows:
            bill.items.append(BillItem(bill_id=bill.id, fee_item_id=row["feeItemId"], fee_item_code=row["feeItemCode"], fee_item_name=row["feeItemName"], amount_fen=row["amountFen"], remark=row.get("remark", "")))
        created += 1
    batch.status = "COMMITTED"
    notification_service.send(db, "BILLS_IMPORTED", "role:ADMIN", {"batchId": batch.id, "created": created})
    db.commit()
    return response({"batchId": batch.id, "createdBills": created, "skippedBills": skipped})


@app.get("/api/admin/payments")
def admin_payments(admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    items = db.scalars(select(Payment).options(selectinload(Payment.bill)).order_by(Payment.created_at.desc())).all()
    return response([payment_data(item) for item in items])


class PaymentConfirmRequest(BaseModel):
    referenceNo: str = Field(default="", max_length=100)


@app.post("/api/admin/payments/{payment_id}/confirm")
def confirm_manual_payment(payment_id: int, payload: PaymentConfirmRequest, admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    payment = db.scalar(select(Payment).where(Payment.id == payment_id).with_for_update())
    if not payment:
        raise HTTPException(404, "缴费记录不存在")
    bill = db.scalar(select(Bill).where(Bill.id == payment.bill_id).with_for_update())
    if payment.status == "SUCCESS":
        return response({"paymentId": payment.id, "status": payment.status, "idempotent": True}, "该笔缴费已经确认")
    if payment.provider != "MANUAL" or payment.status != "PENDING":
        raise HTTPException(409, "当前缴费状态不能人工确认")
    now = datetime.utcnow()
    payment.status, payment.paid_at = "SUCCESS", now
    payment.provider_transaction_id = payload.referenceNo.strip() or f"MANUAL-{payment.payment_no}"
    bill.status, bill.paid_amount_fen, bill.paid_at = "PAID", bill.total_amount_fen, now
    notification_service.send(db, "PAYMENT_CONFIRMED", f"house:{bill.house_id}", {"paymentId": payment.id})
    db.commit()
    return response({"paymentId": payment.id, "status": payment.status})


@app.get("/api/admin/receipts")
def admin_receipts(admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    return response([receipt_data(item) for item in db.scalars(select(Receipt).order_by(Receipt.created_at.desc())).all()])


class StatusPayload(BaseModel):
    status: str = Field(min_length=2, max_length=32)


@app.put("/api/admin/receipts/{receipt_id}/status")
def update_receipt_status(receipt_id: int, payload: StatusPayload, admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    if payload.status not in ("PENDING", "ISSUED", "REJECTED"):
        raise HTTPException(400, "不支持的收据状态")
    item = db.get(Receipt, receipt_id)
    if not item:
        raise HTTPException(404, "收据申请不存在")
    item.status = payload.status
    db.commit()
    return response({"id": item.id, "status": item.status})


@app.get("/api/admin/repairs")
def admin_repairs(admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    items = db.scalars(select(RepairOrder).options(selectinload(RepairOrder.house).selectinload(House.building), selectinload(RepairOrder.technician)).order_by(RepairOrder.created_at.desc())).all()
    return response([repair_data(item) for item in items])


class RepairAssignRequest(BaseModel):
    technicianId: int


@app.put("/api/admin/repairs/{repair_id}/assign")
def assign_repair(repair_id: int, payload: RepairAssignRequest, admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    item = db.get(RepairOrder, repair_id)
    technician = db.scalar(select(User).where(User.id == payload.technicianId, User.role == "TECHNICIAN", User.enabled.is_(True)))
    if not item:
        raise HTTPException(404, "工单不存在")
    if not technician:
        raise HTTPException(404, "维修师傅不存在或已停用")
    item.technician_id = technician.id
    notification_service.send(db, "REPAIR_ASSIGNED", technician.username, {"orderNo": item.order_no})
    db.commit()
    return response({"id": item.id, "technicianId": technician.id})


class MaintenancePayload(BaseModel):
    buildingId: int
    elevatorName: str = Field(min_length=1, max_length=64)
    maintenanceDate: date
    title: str = Field(min_length=2, max_length=180)
    pdfUrl: str = Field(default="", max_length=255)
    fileSize: int = Field(default=0, ge=0)


@app.post("/api/admin/maintenance-records")
def create_maintenance(payload: MaintenancePayload, admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    if not db.get(Building, payload.buildingId):
        raise HTTPException(404, "楼栋不存在")
    item = MaintenanceRecord(building_id=payload.buildingId, elevator_name=payload.elevatorName, maintenance_date=payload.maintenanceDate, title=payload.title, pdf_url=payload.pdfUrl, file_size=payload.fileSize, uploader_name=admin.display_name)
    db.add(item)
    db.commit()
    return response({"id": item.id})


@app.delete("/api/admin/maintenance-records/{record_id}")
def delete_maintenance(record_id: int, admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    item = db.get(MaintenanceRecord, record_id)
    if not item:
        raise HTTPException(404, "维保记录不存在")
    db.delete(item)
    db.commit()
    return response({"id": record_id}, "维保记录已删除")


@app.get("/api/admin/renovations")
def admin_renovations(admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    return response([renovation_data(item) for item in db.scalars(select(Renovation).order_by(Renovation.created_at.desc())).all()])


@app.put("/api/admin/renovations/{renovation_id}/status")
def update_renovation_status(renovation_id: int, payload: StatusPayload, admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    if payload.status not in ("DEPOSIT_UNPAID", "DEPOSIT_PAID", "COMMITMENT_SIGNED", "COMPLETED", "REJECTED"):
        raise HTTPException(400, "不支持的装修登记状态")
    item = db.get(Renovation, renovation_id)
    if not item:
        raise HTTPException(404, "装修登记不存在")
    item.status = payload.status
    db.commit()
    return response({"id": item.id, "status": item.status})


class NoticePayload(BaseModel):
    title: str = Field(min_length=2, max_length=180)
    category: str = Field(min_length=2, max_length=32)
    content: str = Field(min_length=2)
    pinned: bool = False


@app.post("/api/admin/notices")
def create_notice(payload: NoticePayload, admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    item = Notice(title=payload.title, category=payload.category, content=payload.content, pinned=payload.pinned)
    db.add(item)
    db.flush()
    notification_service.send(db, "NOTICE_PUBLISHED", "all", {"noticeId": item.id})
    db.commit()
    return response({"id": item.id})


@app.put("/api/admin/notices/{notice_id}")
def update_notice(notice_id: int, payload: NoticePayload, admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    item = db.get(Notice, notice_id)
    if not item:
        raise HTTPException(404, "公告不存在")
    item.title, item.category, item.content, item.pinned = payload.title, payload.category, payload.content, payload.pinned
    db.commit()
    return response({"id": item.id})


@app.delete("/api/admin/notices/{notice_id}")
def delete_notice(notice_id: int, admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    item = db.get(Notice, notice_id)
    if not item:
        raise HTTPException(404, "公告不存在")
    db.delete(item)
    db.commit()
    return response({"id": notice_id}, "公告已删除")


@app.get("/api/admin/users")
def admin_users(admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    return response([session_data(db, item) for item in db.scalars(select(User).order_by(User.id)).all()])


class UserPayload(BaseModel):
    username: str = Field(min_length=2, max_length=64)
    displayName: str = Field(min_length=1, max_length=64)
    role: str
    phone: str = Field(default="", max_length=24)
    buildingId: int | None = None
    houseId: int | None = None
    password: str | None = Field(default=None, min_length=8, max_length=128)
    enabled: bool = True


def apply_user_payload(item: User, payload: UserPayload) -> None:
    if payload.role not in ("OWNER", "TECHNICIAN", "ADMIN"):
        raise HTTPException(400, "不支持的用户角色")
    item.username, item.display_name, item.role, item.phone = payload.username, payload.displayName, payload.role, payload.phone
    item.building_id, item.house_id, item.enabled = payload.buildingId, payload.houseId, payload.enabled
    if payload.password:
        item.password_hash = hash_password(payload.password)


@app.post("/api/admin/users")
def create_user(payload: UserPayload, admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    if db.scalar(select(User.id).where(User.username == payload.username)):
        raise HTTPException(409, "用户名已存在")
    item = User(username=payload.username, display_name=payload.displayName, role=payload.role)
    apply_user_payload(item, payload)
    db.add(item)
    db.commit()
    return response({"id": item.id})


@app.put("/api/admin/users/{user_id}")
def update_user(user_id: int, payload: UserPayload, admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    item = db.get(User, user_id)
    if not item:
        raise HTTPException(404, "用户不存在")
    if db.scalar(select(User.id).where(User.username == payload.username, User.id != user_id)):
        raise HTTPException(409, "用户名已存在")
    apply_user_payload(item, payload)
    db.commit()
    return response({"id": item.id})


@app.get("/api/admin/buildings")
def admin_buildings(admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    return response([{"id": item.id, "communityName": item.community_name, "name": item.name} for item in db.scalars(select(Building).order_by(Building.id)).all()])


class BuildingPayload(BaseModel):
    communityName: str = Field(min_length=2, max_length=100)
    name: str = Field(min_length=1, max_length=50)


@app.post("/api/admin/buildings")
def create_building(payload: BuildingPayload, admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    if db.scalar(select(Building.id).where(Building.community_name == payload.communityName, Building.name == payload.name)):
        raise HTTPException(409, "楼栋已存在")
    item = Building(community_name=payload.communityName, name=payload.name)
    db.add(item)
    db.commit()
    return response({"id": item.id})


@app.get("/api/admin/houses")
def admin_houses(admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    items = db.scalars(select(House).options(selectinload(House.building)).order_by(House.id)).all()
    return response([{"id": item.id, "buildingId": item.building_id, "buildingName": f"{item.building.community_name}{item.building.name}", "roomNo": item.room_no, "residentCode": item.resident_code, "ownerName": item.owner_name} for item in items])


class HousePayload(BaseModel):
    buildingId: int
    roomNo: str = Field(min_length=1, max_length=32)
    residentCode: str = Field(min_length=2, max_length=64)
    ownerName: str = Field(min_length=1, max_length=64)


@app.post("/api/admin/houses")
def create_house(payload: HousePayload, admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    if not db.get(Building, payload.buildingId):
        raise HTTPException(404, "楼栋不存在")
    if db.scalar(select(House.id).where(or_(House.resident_code == payload.residentCode, (House.building_id == payload.buildingId) & (House.room_no == payload.roomNo)))):
        raise HTTPException(409, "住户编码或楼栋房号已存在")
    item = House(building_id=payload.buildingId, room_no=payload.roomNo, resident_code=payload.residentCode, owner_name=payload.ownerName)
    db.add(item)
    db.commit()
    return response({"id": item.id})


@app.get("/api/admin/notification-outbox")
def admin_notification_outbox(admin: User = Depends(require_roles("ADMIN")), db: Session = Depends(get_db)):
    items = db.scalars(select(NotificationOutbox).order_by(NotificationOutbox.created_at.desc())).all()
    return response([{"id": item.id, "eventType": item.event_type, "recipient": item.recipient, "payload": item.payload, "status": item.status, "createdAt": item.created_at.isoformat()} for item in items])
